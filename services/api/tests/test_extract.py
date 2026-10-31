import io
import json
from pathlib import Path

import pytest
from docx import Document

from legal_api.files.extract import UnsupportedFile, text_of

TR = json.loads((Path(__file__).parent / "fixtures" / "tr" / "test_strings.json").read_text(encoding="utf-8"))


def test_docx_and_text_are_read():
    doc = Document()
    doc.add_paragraph("DAVACI : X")
    doc.add_paragraph(TR["extract_1"])
    buf = io.BytesIO()
    doc.save(buf)
    assert text_of("a.DOCX", buf.getvalue()) == TR["extract_2"]
    assert text_of("a.txt", b"metin") == "metin"


def test_other_types_are_rejected():
    with pytest.raises(UnsupportedFile):
        text_of("a.exe", b"x")
