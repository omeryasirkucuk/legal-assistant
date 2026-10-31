import io
import zipfile

from docx import Document

from legal_api.drafting.export import render
from legal_api.drafting.formats import choose, counted

DRAFT = ("Not.\n---\n**İSTANBUL 1. SULH HUKUK MAHKEMESİ SAYIN HÂKİMLİĞİNE**\n"
         "**DAVACI :** [Ad]\nKira ödenmedi [[E1]]; **30 gün** süre verildi.\n---\n- kontrol")


def test_format_choice():
    assert choose(None, {}) == "docx"
    assert choose("pdf", {"udf": 5}) == "pdf"
    assert choose(None, {"udf": 2, "pdf": 1}) == "udf"
    assert choose(None, {"udf": 2, "pdf": 2}) == "docx"
    assert counted({"pdf": 1}, "pdf") == {"pdf": 2}
    assert counted({"pdf": 1}, None) == {"pdf": 1}


def test_docx_has_only_the_petition():
    doc = Document(io.BytesIO(render(DRAFT, "docx")))
    lines = [p.text for p in doc.paragraphs]
    assert lines == ["İSTANBUL 1. SULH HUKUK MAHKEMESİ SAYIN HÂKİMLİĞİNE", "DAVACI : [Ad]",
                     "Kira ödenmedi; 30 gün süre verildi."]
    assert doc.paragraphs[0].alignment == 1


def test_udf_offsets_point_into_the_text():
    data = render(DRAFT, "udf")
    xml = zipfile.ZipFile(io.BytesIO(data)).read("content.xml").decode()
    body = xml.split("<![CDATA[")[1].split("]]>")[0]
    assert body.startswith("İSTANBUL") and "[[E1]]" not in body
    import re
    for start, length in re.findall(r'startOffset="(\d+)" length="(\d+)"', xml):
        assert int(start) + int(length) <= len(body)
    assert 'format_id="1.8"' in xml


def test_pdf_is_written():
    assert render(DRAFT, "pdf").startswith(b"%PDF")
