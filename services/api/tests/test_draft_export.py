import io
import json
from pathlib import Path

from docx import Document

from legal_api.drafting.adjust import markers
from legal_api.drafting.export import to_docx

TR = json.loads((Path(__file__).parent / "fixtures" / "tr" / "test_strings.json").read_text(encoding="utf-8"))


def test_docx_has_no_markers_and_bold_headings():
    text = TR["draft_export_1"]
    doc = Document(io.BytesIO(to_docx(text)))
    lines = [p.text for p in doc.paragraphs]
    assert lines == [TR["draft_export_2"], "DAVACI: [Ad]", TR["draft_export_3"]]
    assert doc.paragraphs[0].runs[0].bold


def test_markers_are_read():
    assert markers("a [[E1]] b [[E12]] [[E1]]") == {"[[E1]]", "[[E12]]"}


def test_a_request_line_above_the_court_is_set_right_and_the_court_centred():
    from legal_api.drafting.export import blocks
    text = (TR["draft_export_4"])
    aligns = [(b.text, b.align) for b in blocks(text)]
    assert aligns[0] == (TR["draft_export_5"], "right")
    assert aligns[1][1] == "centre" and aligns[-2] == (TR["draft_export_6"], "justify")
