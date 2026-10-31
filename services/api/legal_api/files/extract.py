"""Plain text of a document, for reading the lawyer's samples and attachments.

Read: Word (.docx), PDF, UYAP (.udf), OpenDocument text (.odt), RTF, Excel (.xlsx), e-mail
(.eml), web pages, and plain text formats. Images are not read here; they are shown to the
model as images. Anything else is stored but not read.
"""

from __future__ import annotations

import email
import html
import io
import re
import zipfile
from email import policy

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_BYTES = 10 * 1024 * 1024
TEXT_SUFFIXES = (".txt", ".md", ".csv", ".json", ".xml")
IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".gif", ".webp")


class UnsupportedFile(ValueError):
    pass


def _xml_text(xml: str) -> str:
    xml = re.sub(r"<text:(?:p|h)\b[^>]*>", "\n", xml)
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def _rtf_text(rtf: str) -> str:
    rtf = re.sub(r"\\par[d]?\b", "\n", rtf)
    rtf = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1254"),
                 rtf)
    rtf = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", rtf)
    return rtf


def _xlsx_text(data: bytes) -> str:
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in book.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append("\t".join(cells))
    return "\n".join(lines)


def _eml_text(data: bytes) -> str:
    msg = email.message_from_bytes(data, policy=policy.default)
    head = "\n".join(f"{k}: {msg[k]}" for k in ("From", "To", "Date", "Subject") if msg[k])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body and body.get_content_type() == "text/html":
        text = _xml_text(text)
    return f"{head}\n\n{text}"


def text_of(filename: str, data: bytes) -> str:
    """The document's text; UnsupportedFile for a type that cannot be read."""
    if len(data) > MAX_BYTES:
        raise UnsupportedFile("file too large to read")
    name = filename.lower()
    if name.endswith(".docx"):
        return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if name.endswith(".udf"):
        xml = zipfile.ZipFile(io.BytesIO(data)).read("content.xml").decode("utf-8")
        found = re.search(r"<content><!\[CDATA\[(.*?)\]\]></content>", xml, re.DOTALL)
        return found.group(1) if found else ""
    if name.endswith(".odt"):
        return _xml_text(zipfile.ZipFile(io.BytesIO(data)).read("content.xml").decode("utf-8"))
    if name.endswith(".rtf"):
        return _rtf_text(data.decode("latin-1"))
    if name.endswith(".xlsx"):
        return _xlsx_text(data)
    if name.endswith(".eml"):
        return _eml_text(data)
    if name.endswith((".html", ".htm")):
        return _xml_text(data.decode("utf-8", errors="replace"))
    if name.endswith(TEXT_SUFFIXES):
        return data.decode("utf-8", errors="replace")
    raise UnsupportedFile("unsupported file type")


def readable_text(filename: str, data: bytes) -> str:
    """The text when the type can be read, else empty (the file is still kept)."""
    try:
        return text_of(filename, data)
    except (UnsupportedFile, zipfile.BadZipFile, KeyError, ValueError):
        return ""
