"""A draft as a file: Word (default), PDF or UYAP's UDF.

Only the petition goes into the file, not the note before it or the checklist after it.
Citation markers are removed; headings and **bold** spans stay bold. A draft edited in
the editor brings its own blocks (drafting/rich.py) with italics, underline, sizes and
alignment.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

from docx import Document
from docx.shared import Cm, Pt

from legal_api.drafting.style import petition_part
from legal_api.grounding.normalize import tr_fold
from legal_api.locale import style_terms

FORMATS = ("docx", "pdf", "udf")
MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "udf": "application/octet-stream",
}
FONT = "Times New Roman"
FONT_SIZE = 12
MARGIN_CM = 2.5

_MARKER = re.compile(r"\s*\[\[E\d+\]\]")
_KEY = re.compile(r"\[\[(E\d+)\]\]")
_BOLD = re.compile(r"\*\*(.+?)\*\*")


@dataclass(frozen=True)
class Run:
    """A stretch of text in one style. `marker` is a citation key (E3) the editor shows as a
    chip; files leave markers out."""
    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    size: int | None = None         # points; None is the body size
    marker: str | None = None


@dataclass
class Block:
    """A paragraph of the petition: its runs and its alignment."""
    runs: list[Run]
    align: str = "justify"          # justify | left | centre | right

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs if r.marker is None)


def _plain(line: str) -> str:
    return _MARKER.sub("", line).replace("**", "").lstrip("#").strip()


def _is_heading(line: str) -> bool:
    text = _plain(line)
    return bool(text) and (line.lstrip().startswith("#") or
                           (text.upper() == text and any(c.isalpha() for c in text)))


def _is_request_flag(text: str) -> bool:
    """A line above the court flagging a request (an interim measure, legal aid); lawyers set
    it at the top right."""
    return tr_fold(text).endswith(style_terms()["request_flag"] + "dir")


def _runs(line: str, keep_markers: bool) -> list[Run]:
    """A body line's runs: **bold** spans, and citation markers as their own runs when kept."""
    clean = (line if keep_markers else _MARKER.sub("", line)).lstrip("#").strip()
    runs: list[Run] = []
    pos = 0
    for m in _BOLD.finditer(clean):
        runs += _with_markers(clean[pos:m.start()], False)
        runs += _with_markers(m.group(1), True)
        pos = m.end()
    runs += _with_markers(clean[pos:], False)
    return [r for r in runs if r.text or r.marker]


def _with_markers(text: str, bold: bool) -> list[Run]:
    out, pos = [], 0
    for m in _KEY.finditer(text):
        out += [Run(text[pos:m.start()], bold=bold), Run("", marker=m.group(1))]
        pos = m.end()
    out.append(Run(text[pos:], bold=bold))
    return out


def blocks(text: str, keep_markers: bool = False) -> list[Block]:
    """The petition's paragraphs. With `keep_markers` citation markers stay as runs (for the
    editor); files leave them out."""
    out: list[Block] = []
    court_placed = False
    for line in petition_part(text).splitlines():
        if not line.strip():
            continue
        if _is_heading(line):
            plain = _plain(line)
            if not court_placed and _is_request_flag(plain):
                out.append(Block([Run(plain, bold=True)], align="right"))
                continue
            # The court's name heads the petition and is centred.
            out.append(Block([Run(plain, bold=True)], align="justify" if court_placed else "centre"))
            court_placed = True
            continue
        out.append(Block(_runs(line, keep_markers)))
        court_placed = True
    return out


def to_docx(text_or_blocks: str | list[Block], title: str = "") -> bytes:
    doc = Document()
    # Shown as the document's name by Word and file viewers (else empty or "anonymous").
    doc.core_properties.title = title
    doc.core_properties.author = ""
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(MARGIN_CM))
    style = doc.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(FONT_SIZE)
    paragraphs = blocks(text_or_blocks) if isinstance(text_or_blocks, str) else text_or_blocks
    for block in paragraphs:
        p = doc.add_paragraph()
        p.alignment = {"left": 0, "centre": 1, "right": 2}.get(block.align, 3)
        for r in block.runs:
            if r.marker:
                continue
            run = p.add_run(r.text)
            run.bold, run.italic, run.underline = r.bold, r.italic, r.underline
            if r.size:
                run.font.size = Pt(r.size)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render(text_or_blocks: str | list[Block], fmt: str, title: str = "") -> bytes:
    paragraphs = blocks(text_or_blocks) if isinstance(text_or_blocks, str) else text_or_blocks
    if fmt == "pdf":
        from legal_api.drafting.export_pdf import to_pdf
        return to_pdf(paragraphs, title)
    if fmt == "udf":
        from legal_api.drafting.export_udf import to_udf
        return to_udf(paragraphs)
    return to_docx(paragraphs, title)
